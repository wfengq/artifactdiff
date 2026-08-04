import json
from pathlib import Path

import pytest
from docx import Document

from artifactdiff.service import CompareOptions, compare_documents
from tests.factories import make_pdf


def test_compare_real_changed_pdfs_without_visual_rendering(tmp_path: Path) -> None:
    before = make_pdf(
        tmp_path / "before.pdf",
        [[(72, 720, "Revenue 100")], [(72, 720, "Notes")]],
    )
    after = make_pdf(
        tmp_path / "after.pdf",
        [[(200, 700, "Second"), (72, 720, "First"), (72, 680, "Third")]],
    )

    run = compare_documents(
        before,
        after,
        tmp_path / "report",
        options=CompareOptions(visual=False),
    )

    assert run.result.status == "changed"
    assert run.result.summary.total_changes > 0
    assert run.visual_assets == {}
    assert json.loads(run.json_path.read_text(encoding="utf-8"))["schema_version"] == "1.0"


@pytest.mark.parametrize("hidden", [False, True])
def test_compare_real_docx_detects_explicit_false_and_genuine_hidden_changes(
    tmp_path: Path, hidden: bool
) -> None:
    def make_source(path: Path, term: str) -> Path:
        document = Document()
        document.add_heading("Section 1 Terms", level=1)
        paragraph = document.add_paragraph("Value ")
        run = paragraph.add_run(term)
        run.font.hidden = hidden
        document.save(path)
        return path

    before = make_source(tmp_path / "before.docx", "Term A")
    after = make_source(tmp_path / "after.docx", "Term B")

    run = compare_documents(
        before,
        after,
        tmp_path / "report",
        options=CompareOptions(visual=False),
    )

    assert run.result.status == "changed"
    assert run.result.summary.total_changes > 0
