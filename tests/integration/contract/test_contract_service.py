import json
from hashlib import sha256
from pathlib import Path
from typing import Any, cast

import pytest

from artifactdiff.contract import inspect_contract, load_contract
from artifactdiff.contract.models import ContractClause, ContractDocument, EntityKind
from artifactdiff.contract.selectors import (
    ClauseSelector,
    SelectorResolutionStatus,
    resolve_baseline,
    resolve_candidate,
)
from artifactdiff.errors import InputValidationError
from tests.factories import (
    ContractLanguage,
    _margin_image,
    make_contract_docx,
    make_contract_pdf,
)

FACTORIES = {
    "docx": make_contract_docx,
    "pdf": make_contract_pdf,
}
LANGUAGE_KINDS: dict[ContractLanguage, str] = {
    "zh": "zh",
    "en": "en",
    "zh-en": "zh-en",
}
PAYMENT_ENTITY_KINDS = {
    EntityKind.PARTY,
    EntityKind.MONEY,
    EntityKind.DATE,
    EntityKind.DURATION,
    EntityKind.PERCENTAGE,
}


def _payment_clause(contract: ContractDocument) -> ContractClause:
    matches = [
        clause
        for clause in contract.clauses
        if PAYMENT_ENTITY_KINDS <= {entity.kind for entity in clause.entities}
    ]
    assert len(matches) == 1
    return matches[0]


def _selector_for_payment_clause(contract: ContractDocument) -> ClauseSelector:
    payment = _payment_clause(contract)
    anchor = payment.text.splitlines()[-1]
    return ClauseSelector(
        clause_label=payment.label.normalized,
        heading=payment.heading,
        ancestor_path=payment.ancestor_path,
        anchor=anchor,
        baseline_fingerprint=payment.fingerprint,
    )


@pytest.mark.parametrize("language", ["zh", "en", "zh-en"])
@pytest.mark.parametrize("format_name", ["docx", "pdf"])
def test_contract_factories_are_byte_stable(
    tmp_path: Path, format_name: str, language: ContractLanguage
) -> None:
    factory = FACTORIES[format_name]
    first = factory(tmp_path / f"first.{format_name}", language=language)
    second = factory(tmp_path / f"second.{format_name}", language=language)

    assert second.read_bytes() == first.read_bytes()


def test_pdf_factory_uses_self_contained_cjk_font_without_host_candidates(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from reportlab.pdfbase import pdfmetrics

    monkeypatch.delitem(pdfmetrics._fonts, "ArtifactDiffContractCJK", raising=False)
    with monkeypatch.context() as context:
        context.setattr(Path, "is_file", lambda _: False)
        first = make_contract_pdf(tmp_path / "first.pdf", language="zh")
        second = make_contract_pdf(tmp_path / "second.pdf", language="zh")

    assert second.read_bytes() == first.read_bytes()
    first_ir = inspect_contract(first)
    assert inspect_contract(first) == first_ir
    assert cast(dict[str, Any], first_ir["language"])["kind"] == "zh"


def test_inspect_contract_bounds_independent_headings_with_clauses(tmp_path: Path) -> None:
    path = make_contract_pdf(tmp_path / "contract.pdf", language="zh")

    payload = inspect_contract(path, max_clauses=0)

    assert payload["clauses"] == []
    assert payload["independent_headings"] == []
    assert payload["truncated_headings"] is True


def test_inspect_contract_excludes_docx_running_text_from_headings(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "contract.docx"
    document = Document()
    document.add_heading("第一条 付款条件", level=1)
    document.sections[0].header.paragraphs[0].text = "第一章 合同主体"
    document.sections[0].footer.paragraphs[0].text = "二、 服务范围"
    document.save(source)

    payload = inspect_contract(source)

    assert payload["independent_headings"] == ["第一条 付款条件"]


@pytest.mark.parametrize("fill_line", [True, False])
def test_numbered_pdf_fill_blank_is_not_a_heading(tmp_path: Path, fill_line: bool) -> None:
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas

    path = tmp_path / ("form.pdf" if fill_line else "headings.pdf")
    document = canvas.Canvas(str(path), invariant=1)
    document.setFont("Helvetica", 12)
    document.drawString(100, 700, "Section 2 Termination")
    document.drawString(100, 660, "2.3 Other")
    if fill_line:
        line_start = 100 + stringWidth("2.3 Other", "Helvetica", 12)
        document.line(line_start, 660, 500, 660)
    document.save()

    payload = inspect_contract(path)

    assert payload["independent_headings"] == (
        ["Section 2 Termination"]
        if fill_line
        else ["Section 2 Termination", "2.3 Other"]
    )


def test_two_column_pdf_reports_headings_in_both_columns(tmp_path: Path) -> None:
    from reportlab.pdfgen import canvas

    source = tmp_path / "two-column.pdf"
    document = canvas.Canvas(str(source), pagesize=(595, 842), invariant=1)
    document.setFont("Helvetica", 12)
    for index in range(10):
        y = 760 - index * 20
        left = "Section 1 Payment Terms" if index == 0 else f"Left body line {index}."
        right = "Section 2 Confidentiality" if index == 0 else f"Right body line {index}."
        document.drawString(70, y, left)
        document.drawString(335, y, right)
    document.save()

    payload = inspect_contract(source)

    assert payload["independent_headings"] == [
        "Section 1 Payment Terms",
        "Section 2 Confidentiality",
    ]


def test_margin_images_preserve_distinct_chinese_glyph_content() -> None:
    confidential = _margin_image("机密合同").getvalue()
    attachment = _margin_image("价格附件").getvalue()

    assert sha256(confidential).digest() != sha256(attachment).digest()


@pytest.mark.parametrize("language", ["zh", "en", "zh-en"])
@pytest.mark.parametrize("format_name", ["docx", "pdf"])
def test_inspect_contract_is_stable_for_all_language_and_format_fixtures(
    tmp_path: Path, format_name: str, language: ContractLanguage
) -> None:
    source = FACTORIES[format_name](
        tmp_path / f"contract.{format_name}", language=language
    )

    first = inspect_contract(source)
    second = inspect_contract(source)

    assert second == first
    assert cast(dict[str, Any], first["language"])["kind"] == LANGUAGE_KINDS[language]
    assert len(cast(list[object], first["clauses"])) == 4
    assert PAYMENT_ENTITY_KINDS <= {
        EntityKind(item["kind"])
        for item in cast(list[dict[str, Any]], first["entities"])
    }
    region_kinds = {
        item["kind"]
        for item in cast(list[dict[str, Any]], first["protected_regions"])
    }
    assert {"signature", "seal", "attachment"} <= region_kinds
    if format_name == "docx":
        assert {"header", "footer"} <= region_kinds
    else:
        image_features = [
            item
            for item in cast(list[dict[str, Any]], first["features"])
            if item["kind"] == "embedded_image"
        ]
        assert sum(cast(int, item["count"]) for item in image_features) == 2


def test_inspect_contract_returns_bounded_portable_text_minimizing_ir(
    tmp_path: Path,
) -> None:
    source = make_contract_docx(tmp_path / "contract.docx", language="zh")

    full = inspect_contract(source)
    bounded = inspect_contract(source, max_clauses=2)

    assert len(cast(list[object], bounded["clauses"])) == 2
    assert bounded["truncated_clauses"] is True
    assert (
        cast(list[str], bounded["independent_headings"])
        == cast(list[str], full["independent_headings"])[:2]
    )
    assert bounded["truncated_headings"] is True
    bounded_fields = {"clauses", "truncated_clauses", "independent_headings", "truncated_headings"}
    assert {key: value for key, value in bounded.items() if key not in bounded_fields} == {
        key: value for key, value in full.items() if key not in bounded_fields
    }
    serialized = json.dumps(bounded, ensure_ascii=False, sort_keys=True)
    assert "artifactdiff-contract-" not in serialized
    assert "Private Contract Author" not in serialized


@pytest.mark.parametrize("max_clauses", [-1, 1001])
def test_inspect_contract_validates_clause_bound_before_source_work(
    tmp_path: Path, max_clauses: int
) -> None:
    missing = tmp_path / "missing.docx"

    with pytest.raises(
        InputValidationError, match="max_clauses must be between 0 and 1000"
    ):
        inspect_contract(missing, max_clauses=max_clauses)


def test_inspect_contract_accepts_zero_clause_limit(tmp_path: Path) -> None:
    source = make_contract_docx(tmp_path / "contract.docx", language="en")

    result = inspect_contract(source, max_clauses=0)

    assert result["clauses"] == []
    assert result["truncated_clauses"] is True


def test_inspect_contract_honors_absolute_path_requirement(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    make_contract_docx(tmp_path / "contract.docx", language="en")
    monkeypatch.chdir(tmp_path)

    with pytest.raises(InputValidationError, match="Path must be absolute"):
        inspect_contract(Path("contract.docx"), require_absolute=True)


@pytest.mark.parametrize("language", ["zh", "en", "zh-en"])
def test_docx_and_pdf_contracts_resolve_the_derived_payment_selector(
    tmp_path: Path, language: ContractLanguage
) -> None:
    _, docx = load_contract(
        make_contract_docx(tmp_path / "contract.docx", language=language),
        render=False,
        force=False,
        workdir=tmp_path / "docx-work",
    )
    _, pdf = load_contract(
        make_contract_pdf(tmp_path / "contract.pdf", language=language),
        render=False,
        force=False,
        workdir=tmp_path / "pdf-work",
    )
    selector = _selector_for_payment_clause(docx)
    payment = _payment_clause(docx)

    baseline = resolve_baseline(docx, selector)
    candidate = resolve_candidate(pdf, selector)

    assert baseline.status is SelectorResolutionStatus.UNIQUE
    assert baseline.matches[0].clause_id == payment.id
    assert candidate.status is SelectorResolutionStatus.HIGH_CONFIDENCE


@pytest.mark.parametrize("language", ["zh", "en", "zh-en"])
def test_docx_payment_selector_relocates_30_to_45_day_pdf_edit(
    tmp_path: Path, language: ContractLanguage
) -> None:
    _, baseline = load_contract(
        make_contract_docx(
            tmp_path / "contract.docx", language=language, payment_days=30
        ),
        render=False,
        force=False,
        workdir=tmp_path / "docx-work",
    )
    _, candidate = load_contract(
        make_contract_pdf(
            tmp_path / "contract.pdf", language=language, payment_days=45
        ),
        render=False,
        force=False,
        workdir=tmp_path / "pdf-work",
    )
    selector = _selector_for_payment_clause(baseline)

    baseline_resolution = resolve_baseline(baseline, selector)
    candidate_resolution = resolve_candidate(candidate, selector)

    assert baseline_resolution.status is SelectorResolutionStatus.UNIQUE
    assert candidate_resolution.status is SelectorResolutionStatus.HIGH_CONFIDENCE
    assert candidate_resolution.matches[0].clause_id == _payment_clause(candidate).id


def test_derived_payment_selector_fails_closed_for_ambiguous_candidate(
    tmp_path: Path,
) -> None:
    _, baseline = load_contract(
        make_contract_docx(tmp_path / "contract.docx", language="en"),
        render=False,
        force=False,
        workdir=tmp_path / "docx-work",
    )
    _, candidate = load_contract(
        make_contract_pdf(tmp_path / "contract.pdf", language="en"),
        render=False,
        force=False,
        workdir=tmp_path / "pdf-work",
    )
    payment = _payment_clause(candidate)
    duplicate = payment.model_copy(update={"id": f"{payment.id}-duplicate"})
    ambiguous = candidate.model_copy(update={"clauses": [*candidate.clauses, duplicate]})

    resolution = resolve_candidate(ambiguous, _selector_for_payment_clause(baseline))

    assert resolution.status is SelectorResolutionStatus.AMBIGUOUS


def test_load_contract_keeps_repeated_entity_ids_unique_and_cross_format_stable(
    tmp_path: Path,
) -> None:
    from docx import Document
    from reportlab.pdfgen import canvas

    heading = "Section 1 Repeated values"
    lines = [
        "Party A: Acme Ltd;",
        "Party A: Acme Ltd;",
        "RMB 100 and",
        "RMB 100",
    ]
    docx_path = tmp_path / "repeated.docx"
    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph("\n".join(lines))
    document.save(str(docx_path))
    pdf_path = tmp_path / "repeated.pdf"
    pdf = canvas.Canvas(str(pdf_path), invariant=1)
    pdf.drawString(72, 720, heading)
    for index, line in enumerate(lines):
        pdf.drawString(72, 690 - index * 30, line)
    pdf.save()

    docx_snapshot, docx_contract = load_contract(
        docx_path, render=False, force=False, workdir=tmp_path / "docx-work"
    )
    pdf_snapshot, pdf_contract = load_contract(
        pdf_path, render=False, force=False, workdir=tmp_path / "pdf-work"
    )

    assert len(docx_snapshot.blocks) == 2
    assert len(pdf_snapshot.blocks) == 5
    assert docx_contract.clauses[0].id == pdf_contract.clauses[0].id
    assert len(docx_contract.entities) == 6
    assert len({item.id for item in docx_contract.entities}) == 6
    assert len({item.id for item in pdf_contract.entities}) == 6
    assert [
        (item.kind, item.normalized_value, item.id) for item in pdf_contract.entities
    ] == [
        (item.kind, item.normalized_value, item.id) for item in docx_contract.entities
    ]
    assert len({item.evidence[0].block_id for item in docx_contract.entities}) == 1
    assert len({item.evidence[0].block_id for item in pdf_contract.entities}) == 4
