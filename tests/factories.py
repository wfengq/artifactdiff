from base64 import b64decode
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Literal
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from PIL import Image, ImageDraw, ImageFont

from artifactdiff.models import (
    BlockRef,
    ComparisonResult,
    ComparisonSummary,
    ContentType,
    SemanticChange,
    SourceDescriptor,
    VisualPageChange,
)
from artifactdiff.visual import VisualAssets

SAMPLE_PDF = "JVBERi0xLjMKJZOMi54gUmVwb3J0TGFiIEdlbmVyYXRlZCBQREYgZG9jdW1lbnQgKG9wZW5zb3VyY2UpCjEgMCBvYmoKPDwKL0YxIDIgMCBSCj4+CmVuZG9iagoyIDAgb2JqCjw8Ci9CYXNlRm9udCAvSGVsdmV0aWNhIC9FbmNvZGluZyAvV2luQW5zaUVuY29kaW5nIC9OYW1lIC9GMSAvU3VidHlwZSAvVHlwZTEgL1R5cGUgL0ZvbnQKPj4KZW5kb2JqCjMgMCBvYmoKPDwKL0NvbnRlbnRzIDggMCBSIC9NZWRpYUJveCBbIDAgMCA2MTIgNzkyIF0gL1BhcmVudCA3IDAgUiAvUmVzb3VyY2VzIDw8Ci9Gb250IDEgMCBSIC9Qcm9jU2V0IFsgL1BERiAvVGV4dCAvSW1hZ2VCIC9JbWFnZUMgL0ltYWdlSSBdCj4+IC9Sb3RhdGUgMCAvVHJhbnMgPDwKCj4+IAogIC9UeXBlIC9QYWdlCj4+CmVuZG9iago0IDAgb2JqCjw8Ci9Db250ZW50cyA5IDAgUiAvTWVkaWFCb3ggWyAwIDAgNjEyIDc5MiBdIC9QYXJlbnQgNyAwIFIgL1Jlc291cmNlcyA8PAovRm9udCAxIDAgUiAvUHJvY1NldCBbIC9QREYgL1RleHQgL0ltYWdlQiAvSW1hZ2VDIC9JbWFnZUkgXQo+PiAvUm90YXRlIDAgL1RyYW5zIDw8Cgo+PiAKICAvVHlwZSAvUGFnZQo+PgplbmRvYmoKNSAwIG9iago8PAovUGFnZU1vZGUgL1VzZU5vbmUgL1BhZ2VzIDcgMCBSIC9UeXBlIC9DYXRhbG9nCj4+CmVuZG9iago2IDAgb2JqCjw8Ci9BdXRob3IgKGFub255bW91cykgL0NyZWF0aW9uRGF0ZSAoRDoyMDI2MDgwMzIzMTY0NCswOCcwMCcpIC9DcmVhdG9yIChhbm9ueW1vdXMpIC9LZXl3b3JkcyAoKSAvTW9kRGF0ZSAoRDoyMDI2MDgwMzIzMTY0NCswOCcwMCcpIC9Qcm9kdWNlciAoUmVwb3J0TGFiIFBERiBMaWJyYXJ5IC0gXChvcGVuc291cmNlXCkpIAogIC9TdWJqZWN0ICh1bnNwZWNpZmllZCkgL1RpdGxlICh1bnRpdGxlZCkgL1RyYXBwZWQgL0ZhbHNlCj4+CmVuZG9iago3IDAgb2JqCjw8Ci9Db3VudCAyIC9LaWRzIFsgMyAwIFIgNCAwIFIgXSAvVHlwZSAvUGFnZXMKPj4KZW5kb2JqCjggMCBvYmoKPDwKL0xlbmd0aCA4Ngo+PgpzdHJlYW0KMSAwIDAgMSAwIDAgY20gIEJUIC9GMSAxMiBUZiAxNC40IFRMIEVUCkJUIDEgMCAwIDEgNzIgNzIwIFRtIChSZXZlbnVlIDEwMCkgVGogVCogRVQKIAplbmRzdHJlYW0KZW5kb2JqCjkgMCBvYmoKPDwKL0xlbmd0aCA4MAo+PgpzdHJlYW0KMSAwIDAgMSAwIDAgY20gIEJUIC9GMSAxMiBUZiAxNC40IFRMIEVUCkJUIDEgMCAwIDEgNzIgNzIwIFRtIChOb3RlcykgVGogVCogRVQKIAplbmRzdHJlYW0KZW5kb2JqCnhyZWYKMCAxMAowMDAwMDAwMDAwIDY1NTM1IGYgCjAwMDAwMDAwNjEgMDAwMDAgbiAKMDAwMDAwMDA5MiAwMDAwMCBuIAowMDAwMDAwMTk5IDAwMDAwIG4gCjAwMDAwMDAzOTIgMDAwMDAgbiAKMDAwMDAwMDU4NSAwMDAwMCBuIAowMDAwMDAwNjUzIDAwMDAwIG4gCjAwMDAwMDA5MTQgMDAwMDAgbiAKMDAwMDAwMDk3OSAwMDAwMCBuIAowMDAwMDAxMTE0IDAwMDAwIG4gCnRyYWlsZXIKPDwKL0lEIApbPGU4ZTE2ZTRhMTA2Mzc3NzBkMGU4OTQ2MDFmZTlkZmM1PjxlOGUxNmU0YTEwNjM3NzcwZDBlODk0NjAxZmU5ZGZjNT5dCiUgUmVwb3J0TGFiIGdlbmVyYXRlZCBQREYgZG9jdW1lbnQgLS0gZGlnZXN0IChvcGVuc291cmNlKQoKL0luZm8gNiAwIFIKL1Jvb3QgNSAwIFIKL1NpemUgMTAKPj4Kc3RhcnR4cmVmCjEyNDMKJSVFT0YK"
ORDERED_PDF = "JVBERi0xLjMKJZOMi54gUmVwb3J0TGFiIEdlbmVyYXRlZCBQREYgZG9jdW1lbnQgKG9wZW5zb3VyY2UpCjEgMCBvYmoKPDwKL0YxIDIgMCBSCj4+CmVuZG9iagoyIDAgb2JqCjw8Ci9CYXNlRm9udCAvSGVsdmV0aWNhIC9FbmNvZGluZyAvV2luQW5zaUVuY29kaW5nIC9OYW1lIC9GMSAvU3VidHlwZSAvVHlwZTEgL1R5cGUgL0ZvbnQKPj4KZW5kb2JqCjMgMCBvYmoKPDwKL0NvbnRlbnRzIDcgMCBSIC9NZWRpYUJveCBbIDAgMCA2MTIgNzkyIF0gL1BhcmVudCA2IDAgUiAvUmVzb3VyY2VzIDw8Ci9Gb250IDEgMCBSIC9Qcm9jU2V0IFsgL1BERiAvVGV4dCAvSW1hZ2VCIC9JbWFnZUMgL0ltYWdlSSBdCj4+IC9Sb3RhdGUgMCAvVHJhbnMgPDwKCj4+IAogIC9UeXBlIC9QYWdlCj4+CmVuZG9iago0IDAgb2JqCjw8Ci9QYWdlTW9kZSAvVXNlTm9uZSAvUGFnZXMgNiAwIFIgL1R5cGUgL0NhdGFsb2cKPj4KZW5kb2JqCjUgMCBvYmoKPDwKL0F1dGhvciAoYW5vbnltb3VzKSAvQ3JlYXRpb25EYXRlIChEOjIwMjYwODAzMjMxNjQ0KzA4JzAwJykgL0NyZWF0b3IgKGFub255bW91cykgL0tleXdvcmRzICgpIC9Nb2REYXRlIChEOjIwMjYwODAzMjMxNjQ0KzA4JzAwJykgL1Byb2R1Y2VyIChSZXBvcnRMYWIgUERGIExpYnJhcnkgLSBcKG9wZW5zb3VyY2VcKSkgCiAgL1N1YmplY3QgKHVuc3BlY2lmaWVkKSAvVGl0bGUgKHVudGl0bGVkKSAvVHJhcHBlZCAvRmFsc2UKPj4KZW5kb2JqCjYgMCBvYmoKPDwKL0NvdW50IDEgL0tpZHMgWyAzIDAgUiBdIC9UeXBlIC9QYWdlcwo+PgplbmRvYmoKNyAwIG9iago8PAovTGVuZ3RoIDE1OAo+PgpzdHJlYW0KMSAwIDAgMSAwIDAgY20gIEJUIC9GMSAxMiBUZiAxNC40IFRMIEVUCkJUIDEgMCAwIDEgMjAwIDcwMCBUbSAoU2Vjb25kKSBUaiBUKiBFVApCVCAxIDAgMCAxIDcyIDcyMCBUbSAoRmlyc3QpIFRqIFQqIEVUCkJUIDEgMCAwIDEgNzIgNjgwIFRtIChUaGlyZCkgVGogVCogRVQKIAplbmRzdHJlYW0KZW5kb2JqCnhyZWYKMCA4CjAwMDAwMDAwMDAgNjU1MzUgZiAKMDAwMDAwMDA2MSAwMDAwMCBuIAowMDAwMDAwMDkyIDAwMDAwIG4gCjAwMDAwMDAxOTkgMDAwMDAgbiAKMDAwMDAwMDM5MiAwMDAwMCBuIAowMDAwMDAwNDYwIDAwMDAwIG4gCjAwMDAwMDA3MjEgMDAwMDAgbiAKMDAwMDAwMDc4MCAwMDAwMCBuIAp0cmFpbGVyCjw8Ci9JRCAKWzxiZGU4NjdjOGQwOWZlNmEzN2M1ZmYzYTg2MWM3Mzk5Mj48YmRlODY3YzhkMDlmZTZhMzdjNWZmM2E4NjFjNzM5OTI+XQolIFJlcG9ydExhYiBnZW5lcmF0ZWQgUERGIGRvY3VtZW50IC0tIGRpZ2VzdCAob3BlbnNvdXJjZSkKCi9JbmZvIDUgMCBSCi9Sb290IDQgMCBSCi9TaXplIDgKPj4Kc3RhcnR4cmVmCjk4OAolJUVPRgo="


ContractLanguage = Literal["zh", "en", "zh-en"]
_FIXED_ZIP_TIMESTAMP = (2026, 8, 4, 0, 0, 0)
_FIXTURE_FONT_PATH = (
    Path(__file__).parent / "assets" / "fonts" / "ArtifactDiffContractCJK-Regular.ttf"
)


def contract_lines(language: ContractLanguage, payment_days: int = 30) -> list[str]:
    """Return the shared four-clause body used by contract fixtures."""
    if language == "zh":
        return [
            "第一条 合同主体",
            "本合同由双方于平等基础上订立。",
            "第二条 付款条件",
            f"甲方：上海示例科技有限公司；甲方应于 2026年8月4日 后 {payment_days} 天内支付 RMB 10,000.00，违约金为 5%。",
            "第三条 一般条款",
            "本合同适用中华人民共和国法律。",
            "签字：授权代表",
            "盖章：公司公章",
            "第四条 附件",
            "附件 A：价格表",
        ]
    if language == "en":
        return [
            "Article I Parties",
            "The parties enter this agreement on equal terms.",
            "Article II Payment Terms",
            f"Party A: Example Ltd.; on 2026-08-04, Party A shall pay RMB 10,000.00 within {payment_days} days with a 5% late fee.",
            "Article III General Terms",
            "This agreement is governed by the applicable law.",
            "Signature: Authorized Representative",
            "Company Chop: Official Seal",
            "Article IV Attachment",
            "Attachment A: Price Schedule",
        ]
    if language == "zh-en":
        return [
            "Article I Parties 合同主体",
            "双方依据平等自愿原则订立本合同并共同遵守全部约定。 The parties enter this agreement on equal terms.",
            "Article II Payment Terms 付款条件",
            f"甲方 Party A: Example Ltd.; 甲方应于 2026年8月4日 pay RMB 10,000.00 within {payment_days} 天，违约金 late fee 5%。",
            "Article III General Terms 一般条款",
            "本合同 This agreement is governed by applicable law 适用法律。",
            "签字 Signature: Authorized Representative 授权代表",
            "盖章 Company Chop: Official Seal 公司公章",
            "Article IV Attachment 附件",
            "附件 Attachment A: Price Schedule 价格表",
        ]
    raise ValueError(f"Unsupported contract language: {language}")


def make_contract_docx(path: Path, *, language: ContractLanguage, payment_days: int = 30) -> Path:
    """Create a byte-stable four-clause DOCX contract fixture."""
    document = Document()
    properties = document.core_properties
    properties.author = "Private Contract Author"
    properties.created = datetime(2026, 8, 4, tzinfo=UTC)
    properties.modified = datetime(2026, 8, 4, tzinfo=UTC)
    header, footer = _contract_margins(language)
    document.sections[0].header.paragraphs[0].text = header
    document.sections[0].footer.paragraphs[0].text = footer
    for index, line in enumerate(contract_lines(language, payment_days)):
        if index in {0, 2, 4, 8}:
            document.add_heading(line, level=1)
        else:
            document.add_paragraph(line)
    package = BytesIO()
    document.save(package)
    path.write_bytes(_stable_zip(package.getvalue()))
    return path


def make_contract_pdf(
    path: Path,
    *,
    language: ContractLanguage,
    payment_days: int = 30,
    signature_y_offset: float = 0.0,
    extra_visual_mark: bool = False,
) -> Path:
    """Create a byte-stable four-clause PDF contract fixture."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    font_name = "ArtifactDiffContractCJK"
    try:
        pdfmetrics.getFont(font_name)
    except KeyError:
        pdfmetrics.registerFont(TTFont(font_name, str(_FIXTURE_FONT_PATH)))
    document = canvas.Canvas(str(path), pagesize=letter, invariant=1, pageCompression=1)
    document.setAuthor("Private Contract Author")
    document.setCreator("ArtifactDiff deterministic fixtures")
    document.setTitle("Contract AD-001")
    header, footer = _contract_margins(language)
    document.drawImage(ImageReader(_margin_image(header)), 54, 748, width=504, height=24)
    document.drawImage(ImageReader(_margin_image(footer)), 54, 20, width=504, height=24)
    document.setFont(font_name, 10)
    y = 700
    for line in contract_lines(language, payment_days):
        line_y = y + signature_y_offset if "signature" in line.casefold() or "签字" in line else y
        document.drawString(72, line_y, line)
        y -= 32
    if extra_visual_mark:
        document.rect(500, 400, 18, 18, fill=1, stroke=0)
    document.save()
    return path


def _contract_margins(language: ContractLanguage) -> tuple[str, str]:
    if language == "zh":
        return "机密合同 AD-001", "第 1 页"
    if language == "en":
        return "CONFIDENTIAL CONTRACT AD-001", "PAGE 1"
    return "机密合同 CONFIDENTIAL CONTRACT AD-001", "第 1 页 / PAGE 1"


def _stable_zip(package: bytes) -> bytes:
    source = BytesIO(package)
    destination = BytesIO()
    with (
        ZipFile(source) as archive,
        ZipFile(destination, "w", compression=ZIP_DEFLATED, compresslevel=9) as stable,
    ):
        for member in sorted(archive.infolist(), key=lambda item: item.filename):
            info = ZipInfo(member.filename, _FIXED_ZIP_TIMESTAMP)
            info.compress_type = ZIP_DEFLATED
            info.external_attr = member.external_attr
            stable.writestr(info, archive.read(member.filename))
    return destination.getvalue()


def _margin_image(text: str) -> BytesIO:
    image = Image.new("RGB", (1008, 48), "white")
    font = ImageFont.truetype(str(_FIXTURE_FONT_PATH), size=24)
    ImageDraw.Draw(image).text((8, 6), text, fill="black", font=font)
    output = BytesIO()
    image.save(output, format="PNG", optimize=False)
    output.seek(0)
    return output


def make_pdf(path: Path, pages: list[list[tuple[float, float, str]]]) -> Path:
    fixtures = {
        (("Revenue 100",), ("Notes",)): SAMPLE_PDF,
        (("Second", "First", "Third"),): ORDERED_PDF,
    }
    key = tuple(tuple(text for _, _, text in page) for page in pages)
    path.write_bytes(b64decode(fixtures[key]))
    return path


def make_docx(path: Path, *, heading: str, paragraphs: list[str], rows: list[list[str]]) -> Path:
    document = Document()
    document.add_heading(heading, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    document.save(str(path))
    return path


def sample_changed_result_with_images(
    tmp_path: Path,
) -> tuple[ComparisonResult, dict[str, VisualAssets]]:
    before_image = tmp_path / "before.png"
    after_image = tmp_path / "after.png"
    heatmap_image = tmp_path / "heatmap.png"
    Image.new("RGB", (64, 64), "white").save(before_image)
    Image.new("RGB", (64, 64), "gray").save(after_image)
    Image.new("RGB", (64, 64), "red").save(heatmap_image)
    before = SourceDescriptor(path="before.pdf", sha256="a" * 64, format="pdf", size_bytes=1)
    after = SourceDescriptor(path="after.pdf", sha256="b" * 64, format="pdf", size_bytes=1)
    change = SemanticChange(
        id="change-1",
        kind="modified",
        content_type=ContentType.PDF_TEXT,
        similarity=0.5,
        before=BlockRef(block_id="before-1", ordinal=0, page_index=0, text="Old"),
        after=BlockRef(block_id="after-1", ordinal=0, page_index=0, text="New"),
    )
    visual = VisualPageChange(id="page-1", before_page=1, after_page=1, changed_pixel_ratio=1.0)
    result = ComparisonResult(
        status="changed",
        before=before,
        after=after,
        summary=ComparisonSummary(modified=1, total_changes=1, visual_change_ratio=1.0),
        changes=[change],
        visual_changes=[visual],
    )
    assets = VisualAssets(before_image, after_image, heatmap_image)
    return result, {"page-1": assets}
