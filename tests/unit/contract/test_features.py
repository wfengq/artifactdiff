from __future__ import annotations

import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.shared import Inches
from lxml import etree
from PIL import Image

from artifactdiff.contract.analyzer import analyze_contract
from artifactdiff.contract.features import inspect_docx_features
from artifactdiff.formats.docx import DocxAdapter

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"


def _png_bytes(color: str) -> bytes:
    stream = BytesIO()
    Image.new("RGB", (8, 8), color).save(stream, format="PNG")
    return stream.getvalue()


def _rewrite_package(path: Path, replacements: dict[str, bytes]) -> None:
    with ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    entries.update(replacements)
    with ZipFile(path, "w", ZIP_DEFLATED) as target:
        for name in sorted(entries):
            target.writestr(name, entries[name])


def make_featured_docx(
    path: Path,
    *,
    comment: str = "Internal comment alpha",
    hidden: str = "Hidden term alpha",
    external_url: str = "https://example.test/alpha",
    image_color: str = "red",
    modified: str = "2026-08-04T00:00:00Z",
) -> Path:
    document = Document()
    document.core_properties.author = "Fixture Author"
    document.core_properties.last_modified_by = "Fixture Editor"
    document.core_properties.created = datetime(2026, 8, 1, tzinfo=UTC)
    document.core_properties.modified = datetime(2026, 8, 4, tzinfo=UTC)
    paragraph = document.add_paragraph("Signature: Alice ")
    paragraph.add_run().add_picture(BytesIO(_png_bytes(image_color)), width=Inches(0.1))
    document.save(path)

    with ZipFile(path) as package:
        document_xml = etree.fromstring(package.read("word/document.xml"))
        relationships = etree.fromstring(package.read("word/_rels/document.xml.rels"))
        content_types = etree.fromstring(package.read("[Content_Types].xml"))
        core = etree.fromstring(package.read("docProps/core.xml"))

    paragraph_xml = document_xml.find(f".//{{{W}}}p")
    assert paragraph_xml is not None
    hidden_run = etree.fromstring(
        (
            f'<w:r xmlns:w="{W}"><w:rPr><w:vanish/></w:rPr>'
            f"<w:t>{hidden}</w:t></w:r>"
        ).encode()
    )
    paragraph_xml.append(hidden_run)
    revision = etree.fromstring(
        (
            f'<w:ins xmlns:w="{W}" w:id="1"><w:r><w:t>Revision alpha</w:t>'
            "</w:r></w:ins>"
        ).encode()
    )
    paragraph_xml.append(revision)
    comment_reference = etree.fromstring(
        (
            f'<w:r xmlns:w="{W}"><w:commentReference w:id="0"/></w:r>'
        ).encode()
    )
    paragraph_xml.append(comment_reference)

    etree.SubElement(
        relationships,
        f"{{{PKG_REL}}}Relationship",
        Id="rIdExternal",
        Type=f"{R}/hyperlink",
        Target=external_url,
        TargetMode="External",
    )
    etree.SubElement(
        relationships,
        f"{{{PKG_REL}}}Relationship",
        Id="rIdComments",
        Type=f"{R}/comments",
        Target="comments.xml",
    )
    etree.SubElement(
        content_types,
        f"{{{CONTENT_TYPES}}}Override",
        PartName="/word/comments.xml",
        ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
    )

    modified_node = core.find("{http://purl.org/dc/terms/}modified")
    assert modified_node is not None
    modified_node.text = modified
    comments = etree.fromstring(
        (
            f'<w:comments xmlns:w="{W}"><w:comment w:id="0" w:author="Reviewer">'
            f"<w:p><w:r><w:t>{comment}</w:t></w:r></w:p>"
            "</w:comment></w:comments>"
        ).encode()
    )
    _rewrite_package(
        path,
        {
            "word/document.xml": etree.tostring(document_xml, xml_declaration=True),
            "word/_rels/document.xml.rels": etree.tostring(
                relationships, xml_declaration=True
            ),
            "[Content_Types].xml": etree.tostring(content_types, xml_declaration=True),
            "docProps/core.xml": etree.tostring(core, xml_declaration=True),
            "word/comments.xml": etree.tostring(comments, xml_declaration=True),
        },
    )
    return path


def _feature_fingerprints(path: Path, workdir: Path) -> dict[str, list[str]]:
    snapshot = DocxAdapter().load(path, render=False, workdir=workdir)
    contract = analyze_contract(snapshot)
    grouped: dict[str, list[str]] = {}
    for feature in contract.features:
        grouped.setdefault(feature.kind.value, []).append(feature.fingerprint)
    return {kind: sorted(values) for kind, values in sorted(grouped.items())}


def _serialized_features(path: Path, workdir: Path) -> list[object]:
    snapshot = DocxAdapter().load(path, render=False, workdir=workdir)
    return snapshot.metadata["document_features"]


def test_docx_features_detect_comments_revisions_hidden_text_links_and_metadata(
    tmp_path: Path,
) -> None:
    source = make_featured_docx(tmp_path / "features.docx")

    snapshot = DocxAdapter().load(source, render=False, workdir=tmp_path / "work")
    contract = analyze_contract(snapshot)

    assert {item.kind.value for item in contract.features} == {
        "comment",
        "tracked_revision",
        "hidden_text",
        "external_link",
        "embedded_image",
        "metadata",
    }
    report = json.dumps(contract.model_dump(mode="json"), sort_keys=True)
    assert "Internal comment alpha" not in report
    assert "Hidden term alpha" not in report
    assert "https://example.test/alpha" not in report
    assert all(len(item.fingerprint) == 64 for item in contract.features)


def test_docx_feature_json_is_stable_and_each_mutation_is_isolated_by_kind(
    tmp_path: Path,
) -> None:
    baseline = make_featured_docx(tmp_path / "baseline.docx")
    repeated = make_featured_docx(tmp_path / "repeated.docx")
    baseline_fingerprints = _feature_fingerprints(baseline, tmp_path / "baseline-work")

    assert _serialized_features(repeated, tmp_path / "repeated-work") == _serialized_features(
        baseline, tmp_path / "baseline-features-work"
    )

    mutations = {
        "comment": {"comment": "Internal comment beta"},
        "hidden_text": {"hidden": "Hidden term beta"},
        "external_link": {"external_url": "https://other.test/beta"},
        "embedded_image": {"image_color": "blue"},
        "metadata": {"modified": "2026-08-05T00:00:00Z"},
    }
    for expected_kind, options in mutations.items():
        mutated = make_featured_docx(tmp_path / f"{expected_kind}.docx", **options)
        actual = _feature_fingerprints(mutated, tmp_path / f"{expected_kind}-work")
        assert {
            kind for kind in baseline_fingerprints if baseline_fingerprints[kind] != actual[kind]
        } == {expected_kind}


def test_signature_region_includes_same_block_image_fingerprint_and_evidence(
    tmp_path: Path,
) -> None:
    source = make_featured_docx(tmp_path / "signature.docx")
    contract = analyze_contract(
        DocxAdapter().load(source, render=False, workdir=tmp_path / "work")
    )

    image = next(item for item in contract.features if item.kind.value == "embedded_image")
    signature = next(item for item in contract.protected_regions if item.kind.value == "signature")

    assert signature.feature_fingerprints == [image.fingerprint]
    assert image.evidence
    assert image.evidence[0] in signature.evidence


def test_hidden_text_in_table_is_reported_only_as_a_bounded_feature(tmp_path: Path) -> None:
    source = tmp_path / "hidden-table.docx"
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Visible term "
    hidden_run = cell.paragraphs[0].add_run("Secret table term")
    hidden_run.font.hidden = True
    document.save(source)

    snapshot = DocxAdapter().load(source, render=False, workdir=tmp_path / "work")
    contract = analyze_contract(snapshot)

    assert snapshot.blocks[0].text.strip() == "Visible term"
    assert any(item.kind.value == "hidden_text" for item in contract.features)
    assert "Secret table term" not in contract.model_dump_json()


def test_docx_image_after_hidden_only_table_uses_actual_signature_block_evidence(
    tmp_path: Path,
) -> None:
    source = tmp_path / "hidden-table-signature.docx"
    document = Document()
    hidden_cell_run = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run(
        "Hidden-only table"
    )
    hidden_cell_run.font.hidden = True
    signature = document.add_paragraph("Signature: Alice ")
    signature.add_run().add_picture(BytesIO(_png_bytes("red")), width=Inches(0.1))
    document.save(source)

    snapshot = DocxAdapter().load(source, render=False, workdir=tmp_path / "work")
    contract = analyze_contract(snapshot)
    image = next(item for item in contract.features if item.kind.value == "embedded_image")

    assert [block.id for block in snapshot.blocks] == [image.evidence[0].block_id]
    assert contract.protected_regions[0].feature_fingerprints == [image.fingerprint]


def test_docx_table_image_maps_to_containing_table_block(tmp_path: Path) -> None:
    source = tmp_path / "table-image.docx"
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Signature: Alice "
    cell.paragraphs[0].add_run().add_picture(BytesIO(_png_bytes("red")), width=Inches(0.1))
    document.save(source)

    snapshot = DocxAdapter().load(source, render=False, workdir=tmp_path / "work")
    image = next(
        item
        for item in analyze_contract(snapshot).features
        if item.kind.value == "embedded_image"
    )

    assert image.evidence[0].block_id == snapshot.blocks[0].id


def test_docx_feature_inspection_ignores_unrelated_malformed_custom_xml(tmp_path: Path) -> None:
    source = make_featured_docx(tmp_path / "custom-xml.docx")
    expected = [item.model_dump(mode="json") for item in inspect_docx_features(source)]
    _rewrite_package(source, {"customXml/unrelated.xml": b"<malformed"})

    actual = [item.model_dump(mode="json") for item in inspect_docx_features(source)]

    assert actual == expected
