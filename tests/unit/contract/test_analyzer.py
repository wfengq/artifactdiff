from pathlib import Path

import pytest
from docx import Document
from pydantic import JsonValue

import artifactdiff.formats.docx as docx_module
from artifactdiff.contract.analyzer import analyze_contract
from artifactdiff.contract.models import EntityKind, EvidenceRef, ProtectedRegionKind
from artifactdiff.formats.docx import DocxAdapter
from artifactdiff.formats.pdf import PdfAdapter
from artifactdiff.models import (
    ContentBlock,
    ContentType,
    DocumentSnapshot,
    PageSnapshot,
    Rect,
)
from artifactdiff.normalize import normalize_text


def _rendered_snapshot(*lines: str) -> DocumentSnapshot:
    blocks = [
        ContentBlock(
            id=f"pdf:0:{ordinal}",
            ordinal=ordinal,
            page_index=0,
            content_type=ContentType.PDF_TEXT,
            text=text,
            normalized_text=normalize_text(text),
            bbox=Rect(x0=10, y0=20 + ordinal * 20, x1=200, y1=35 + ordinal * 20),
        )
        for ordinal, text in enumerate(lines)
    ]
    return DocumentSnapshot(
        source_path="rendered.pdf",
        format="pdf",
        sha256="f" * 64,
        size_bytes=1,
        pages=[
            PageSnapshot(
                index=0,
                width=612,
                height=792,
                text="\n".join(lines),
                normalized_text=normalize_text("\n".join(lines)),
                blocks=blocks,
            )
        ],
        blocks=blocks,
    )


@pytest.fixture
def snapshot_with_nested_clauses() -> DocumentSnapshot:
    blocks = [
        ContentBlock(
            id="block-1",
            ordinal=0,
            content_type=ContentType.HEADING,
            text="\u7b2c\u56db\u6761 \u4ed8\u6b3e\u6761\u4ef6",
            normalized_text=normalize_text("\u7b2c\u56db\u6761 \u4ed8\u6b3e\u6761\u4ef6"),
        ),
        ContentBlock(
            id="block-2",
            ordinal=1,
            content_type=ContentType.PARAGRAPH,
            text="\uff08\u4e00\uff09\u4ed8\u6b3e\u65f6\u95f4",
            normalized_text=normalize_text("\uff08\u4e00\uff09\u4ed8\u6b3e\u65f6\u95f4"),
        ),
        ContentBlock(
            id="block-3",
            ordinal=2,
            content_type=ContentType.PARAGRAPH,
            text="\u4e59\u65b9\u5e94\u5728 45 \u5929\u5185\u652f\u4ed8 RMB 10,000.00\u3002",
            normalized_text=normalize_text("\u4e59\u65b9\u5e94\u5728 45 \u5929\u5185\u652f\u4ed8 RMB 10,000.00\u3002"),
        ),
        ContentBlock(
            id="block-4",
            ordinal=3,
            content_type=ContentType.PARAGRAPH,
            text="4.2 \u4ed8\u6b3e\u91d1\u989d",
            normalized_text=normalize_text("4.2 \u4ed8\u6b3e\u91d1\u989d"),
        ),
    ]
    return DocumentSnapshot(
        source_path="contract.docx",
        format="docx",
        sha256="a" * 64,
        size_bytes=1,
        blocks=blocks,
    )


def test_analyzer_builds_parent_child_hierarchy(
    snapshot_with_nested_clauses: DocumentSnapshot,
) -> None:
    contract = analyze_contract(snapshot_with_nested_clauses)

    assert [clause.label.printed for clause in contract.clauses] == ["\u7b2c\u56db\u6761", "\uff08\u4e00\uff09", "4.2"]
    assert contract.clauses[1].parent_id == contract.clauses[0].id
    assert contract.clauses[1].ancestor_path == ("\u4ed8\u6b3e\u6761\u4ef6",)
    assert contract.clauses[1].id in contract.clauses[0].child_ids
    assert contract.clauses[1].entities[0].evidence == [EvidenceRef(block_id="block-3")]


def test_analyzer_preserves_tables_and_classifies_anchored_protected_regions() -> None:
    blocks = [
        ContentBlock(id="body", ordinal=1, content_type=ContentType.PARAGRAPH, text="Preface", normalized_text="preface"),
        ContentBlock(id="heading", ordinal=2, content_type=ContentType.HEADING, text="Section 1 Terms", normalized_text="section 1 terms"),
        ContentBlock(id="table", ordinal=3, content_type=ContentType.TABLE, text="RMB 100", normalized_text="rmb 100"),
        ContentBlock(id="signature", ordinal=4, content_type=ContentType.PARAGRAPH, text="\u7b7e\u5b57\uff1aAlice", normalized_text="\u7b7e\u5b57\uff1aalice"),
        ContentBlock(id="header", ordinal=5, content_type=ContentType.HEADER, text="Confidential", normalized_text="confidential"),
        ContentBlock(id="footer", ordinal=6, content_type=ContentType.FOOTER, text="Page 1", normalized_text="page 1"),
        ContentBlock(id="attachment", ordinal=7, content_type=ContentType.PARAGRAPH, text="Attachment A", normalized_text="attachment a"),
    ]
    snapshot = DocumentSnapshot(
        source_path="contract.pdf", format="pdf", sha256="b" * 64, size_bytes=2, blocks=blocks
    )

    contract = analyze_contract(snapshot)

    assert [evidence.block_id for evidence in contract.tables] == ["table"]
    assert "RMB 100" in contract.clauses[-1].text
    assert [region.kind for region in contract.protected_regions] == [
        ProtectedRegionKind.SIGNATURE,
        ProtectedRegionKind.HEADER,
        ProtectedRegionKind.FOOTER,
        ProtectedRegionKind.ATTACHMENT,
    ]


def test_analyzer_disambiguates_duplicate_base_ids_using_full_clause_text() -> None:
    long_prefix = "x" * 180
    blocks = [
        ContentBlock(
            id="heading-1",
            ordinal=0,
            content_type=ContentType.HEADING,
            text="Section 1 Shared heading",
            normalized_text="section 1 shared heading",
        ),
        ContentBlock(
            id="body-1",
            ordinal=1,
            content_type=ContentType.PARAGRAPH,
            text=f"{long_prefix} first ending RMB 100",
            normalized_text=normalize_text(f"{long_prefix} first ending RMB 100"),
        ),
        ContentBlock(
            id="heading-2",
            ordinal=2,
            content_type=ContentType.HEADING,
            text="Section 1 Shared heading",
            normalized_text="section 1 shared heading",
        ),
        ContentBlock(
            id="body-2",
            ordinal=3,
            content_type=ContentType.PARAGRAPH,
            text=f"{long_prefix} second ending RMB 200",
            normalized_text=normalize_text(f"{long_prefix} second ending RMB 200"),
        ),
    ]
    snapshot = DocumentSnapshot(
        source_path="contract.docx", format="docx", sha256="c" * 64, size_bytes=3, blocks=blocks
    )

    first = analyze_contract(snapshot)
    second = analyze_contract(snapshot)

    assert len({clause.id for clause in first.clauses}) == 2
    assert [clause.id for clause in second.clauses] == [clause.id for clause in first.clauses]
    assert all(
        entity.clause_id == clause.id
        for clause in first.clauses
        for entity in clause.entities
    )


def test_analyzer_keeps_every_repeated_entity_occurrence_unique_and_stable() -> None:
    blocks = [
        ContentBlock(
            id="heading",
            ordinal=0,
            content_type=ContentType.HEADING,
            text="Section 1 Repeated values",
            normalized_text="section 1 repeated values",
        ),
        ContentBlock(
            id="body",
            ordinal=1,
            content_type=ContentType.PARAGRAPH,
            text="Party A: Acme Ltd; Party A: Acme Ltd; RMB 100 and RMB 100",
            normalized_text="party a: acme ltd; party a: acme ltd; rmb 100 and rmb 100",
        ),
    ]
    snapshot = DocumentSnapshot(
        source_path="contract.docx",
        format="docx",
        sha256="e" * 64,
        size_bytes=5,
        blocks=blocks,
    )

    first = analyze_contract(snapshot)
    second = analyze_contract(snapshot)

    assert len(first.entities) == 6
    assert len({item.id for item in first.entities}) == len(first.entities)
    assert [item.id for item in first.clauses[0].entities] == [
        item.id for item in first.entities
    ]
    assert second.model_dump(mode="json") == first.model_dump(mode="json")


def test_analyzer_propagates_docx_rendered_geometry_to_all_evidence_consumers(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "contract.docx"
    document = Document()
    document.add_heading("Section 4 Payment Terms", level=1)
    document.add_paragraph("Party A: Example Ltd.; pay RMB 10,000.00 within 30 days")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "RMB 10,000.00"
    document.add_paragraph("Signature: Alice")
    document.save(str(source))
    rendered = _rendered_snapshot(
        "Section 4 Payment Terms",
        "Party A: Example Ltd.; pay RMB 10,000.00 within 30 days",
        "RMB 10,000.00",
        "Signature: Alice",
    )
    monkeypatch.setattr(docx_module, "find_libreoffice", lambda: tmp_path / "soffice")
    monkeypatch.setattr(
        docx_module,
        "convert_docx_to_pdf",
        lambda source, output_dir, executable: tmp_path / "rendered.pdf",
    )
    monkeypatch.setattr(
        "artifactdiff.formats.docx.PdfAdapter.load", lambda *args, **kwargs: rendered
    )

    snapshot = DocxAdapter().load(source, render=True, workdir=tmp_path / "work")
    contract = analyze_contract(snapshot)

    clause_evidence = contract.clauses[0].evidence[0]
    money = next(entity for entity in contract.entities if entity.kind is EntityKind.MONEY)
    table_evidence = contract.tables[0]
    signature = next(
        region
        for region in contract.protected_regions
        if region.kind is ProtectedRegionKind.SIGNATURE
    )
    assert (clause_evidence.rendered_page_index, clause_evidence.rendered_bbox) == (
        0,
        Rect(x0=10, y0=20, x1=200, y1=35),
    )
    assert (money.evidence[0].rendered_page_index, money.evidence[0].rendered_bbox) == (
        0,
        Rect(x0=10, y0=40, x1=200, y1=55),
    )
    assert (table_evidence.rendered_page_index, table_evidence.rendered_bbox) == (
        0,
        Rect(x0=10, y0=60, x1=200, y1=75),
    )
    assert (
        signature.evidence[0].rendered_page_index,
        signature.evidence[0].rendered_bbox,
    ) == (0, Rect(x0=10, y0=80, x1=200, y1=95))


@pytest.mark.parametrize(
    "invalid_metadata",
    [
        {"rendered_page_index": 0},
        {
            "rendered_page_index": 0,
            "rendered_bbox": {"x0": 1, "y0": 2, "x1": 3},
        },
        {
            "rendered_page_index": "not-a-page",
            "rendered_bbox": {"x0": 1, "y0": 2, "x1": 3, "y1": 4},
        },
    ],
)
def test_analyzer_ignores_partial_or_invalid_rendered_geometry(
    invalid_metadata: dict[str, JsonValue],
) -> None:
    valid_bbox: dict[str, JsonValue] = {"x0": 5, "y0": 6, "x1": 7, "y1": 8}
    blocks = [
        ContentBlock(
            id="heading",
            ordinal=0,
            content_type=ContentType.HEADING,
            text="Section 4 Payment",
            normalized_text="section 4 payment",
            metadata={"rendered_page_index": 1, "rendered_bbox": valid_bbox},
        ),
        ContentBlock(
            id="body",
            ordinal=1,
            content_type=ContentType.PARAGRAPH,
            text="within 30 days",
            normalized_text="within 30 days",
            metadata=invalid_metadata,
        ),
    ]
    snapshot = DocumentSnapshot(
        source_path="contract.docx",
        format="docx",
        sha256="d" * 64,
        size_bytes=4,
        blocks=blocks,
    )

    contract = analyze_contract(snapshot)

    assert contract.clauses[0].evidence[0].rendered_page_index == 1
    assert contract.clauses[0].evidence[0].rendered_bbox == Rect(
        x0=5, y0=6, x1=7, y1=8
    )
    assert contract.clauses[0].evidence[1].rendered_page_index is None
    assert contract.clauses[0].evidence[1].rendered_bbox is None


def test_analyzer_preserves_real_multipage_pdf_clause_and_entity_order(
    tmp_path: Path,
) -> None:
    from reportlab.pdfgen import canvas

    source = tmp_path / "multipage-contract.pdf"
    document = canvas.Canvas(str(source), invariant=1)
    document.drawString(72, 720, "Article I Parties")
    document.drawString(72, 690, "Party A: Alpha Ltd.")
    document.showPage()
    document.drawString(72, 720, "Article II Payment Terms")
    document.drawString(72, 690, "Pay RMB 10,000.00 within 30 days")
    document.save()
    snapshot = PdfAdapter().load(source, render=False, workdir=tmp_path / "work")

    first = analyze_contract(snapshot)
    second = analyze_contract(snapshot)

    assert [(block.page_index, block.ordinal) for block in snapshot.blocks] == [
        (0, 0),
        (0, 1),
        (1, 0),
        (1, 1),
    ]
    assert [clause.heading for clause in first.clauses] == ["Parties", "Payment Terms"]
    assert "Alpha Ltd." in first.clauses[0].text
    assert "RMB 10,000.00" not in first.clauses[0].text
    assert "RMB 10,000.00" in first.clauses[1].text
    assert any(
        entity.kind is EntityKind.PARTY for entity in first.clauses[0].entities
    )
    assert any(
        entity.kind is EntityKind.MONEY for entity in first.clauses[1].entities
    )
    assert second.model_dump(mode="json") == first.model_dump(mode="json")
