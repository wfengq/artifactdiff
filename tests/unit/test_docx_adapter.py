from pathlib import Path

import pytest
from docx import Document

import artifactdiff.formats.docx as docx_module
from artifactdiff.errors import InputValidationError, RenderUnavailableError
from artifactdiff.formats.docx import DocxAdapter
from artifactdiff.models import (
    ContentBlock,
    ContentType,
    DocumentSnapshot,
    PageSnapshot,
    Rect,
)
from artifactdiff.normalize import normalize_text
from tests.factories import make_docx, make_pdf


def _rendered_snapshot(*lines: str) -> DocumentSnapshot:
    blocks = [
        ContentBlock(
            id=f"pdf:0:{ordinal}",
            ordinal=ordinal,
            page_index=0,
            content_type=ContentType.PDF_TEXT,
            text=text,
            normalized_text=normalize_text(text),
            bbox=Rect(x0=10, y0=20 + ordinal * 20, x1=200, y1=35 + ordinal * 20),
        )
        for ordinal, text in enumerate(lines)
    ]
    return DocumentSnapshot(
        source_path="rendered.pdf",
        format="pdf",
        sha256="f" * 64,
        size_bytes=1,
        pages=[
            PageSnapshot(
                index=0,
                width=612,
                height=792,
                text="\n".join(lines),
                normalized_text=normalize_text("\n".join(lines)),
                blocks=blocks,
            )
        ],
        blocks=blocks,
    )


def test_docx_adapter_preserves_heading_paragraph_table_order(tmp_path: Path) -> None:
    source = make_docx(
        tmp_path / "sample.docx",
        heading="Q2 Results",
        paragraphs=["Revenue increased."],
        rows=[["Region", "Total"], ["APAC", "100"]],
    )

    snapshot = DocxAdapter().load(source, render=False, workdir=tmp_path / "work")

    assert [block.content_type for block in snapshot.blocks[:3]] == [
        ContentType.HEADING,
        ContentType.PARAGRAPH,
        ContentType.TABLE,
    ]
    assert snapshot.blocks[0].metadata["level"] == 1
    assert snapshot.blocks[2].text == "Region\tTotal\nAPAC\t100"
    assert snapshot.blocks[2].metadata == {"rows": 2, "columns": 2}


def test_docx_adapter_serializes_bounded_document_features(tmp_path: Path) -> None:
    source = tmp_path / "metadata.docx"
    document = Document()
    document.core_properties.author = "Private Fixture Author"
    document.add_paragraph("Terms")
    document.save(source)

    snapshot = DocxAdapter().load(source, render=False, workdir=tmp_path / "work")

    facts = snapshot.metadata["document_features"]
    assert isinstance(facts, list)
    assert any(item["kind"] == "metadata" for item in facts)
    assert "Private Fixture Author" not in str(facts)


def test_docx_adapter_preserves_lists_with_stable_dense_block_identity(tmp_path: Path) -> None:
    source = tmp_path / "lists.docx"
    document = Document()
    document.add_paragraph("")
    document.add_paragraph("Review revenue", style="List Bullet")
    document.add_paragraph("Approve forecast", style="List Number")
    document.save(source)

    first = DocxAdapter().load(source, render=False, workdir=tmp_path / "first")
    second = DocxAdapter().load(source, render=False, workdir=tmp_path / "second")

    assert [block.text for block in first.blocks] == ["Review revenue", "Approve forecast"]
    assert [block.content_type for block in first.blocks] == [
        ContentType.PARAGRAPH,
        ContentType.PARAGRAPH,
    ]
    assert [block.metadata for block in first.blocks] == [
        {"list_kind": "bullet", "level": 1},
        {"list_kind": "number", "level": 1},
    ]
    assert [block.ordinal for block in first.blocks] == [0, 1]
    assert [block.id for block in first.blocks] == [block.id for block in second.blocks]


def test_docx_adapter_appends_unique_non_empty_headers_and_footers(tmp_path: Path) -> None:
    source = tmp_path / "sections.docx"
    document = Document()
    document.add_paragraph("Body text")
    document.sections[0].header.paragraphs[0].text = "Quarterly report"
    document.sections[0].footer.paragraphs[0].text = "Confidential"
    document.add_section()
    document.save(source)

    snapshot = DocxAdapter().load(source, render=False, workdir=tmp_path / "work")

    assert [block.content_type for block in snapshot.blocks] == [
        ContentType.PARAGRAPH,
        ContentType.HEADER,
        ContentType.FOOTER,
    ]
    assert [block.text for block in snapshot.blocks] == [
        "Body text",
        "Quarterly report",
        "Confidential",
    ]
    assert [block.ordinal for block in snapshot.blocks] == [0, 1, 2]


def test_docx_adapter_keeps_hidden_header_and_footer_text_out_of_semantic_blocks(
    tmp_path: Path,
) -> None:
    source = tmp_path / "hidden-sections.docx"
    document = Document()
    document.add_paragraph("Body text")
    header = document.sections[0].header.paragraphs[0]
    header.add_run("Visible header")
    hidden_header = header.add_run(" Private header note")
    hidden_header.font.hidden = True
    hidden_footer = document.sections[0].footer.paragraphs[0].add_run(
        "Private footer note"
    )
    hidden_footer.font.hidden = True
    document.save(source)

    snapshot = DocxAdapter().load(source, render=False, workdir=tmp_path / "work")

    assert [block.text for block in snapshot.blocks] == ["Body text", "Visible header"]
    serialized = snapshot.model_dump_json()
    assert "Private header note" not in serialized
    assert "Private footer note" not in serialized
    hidden_parts = {
        item["details"]["part_name"]
        for item in snapshot.metadata["document_features"]
        if item["kind"] == "hidden_text"
    }
    assert any(part.startswith("word/header") for part in hidden_parts)
    assert any(part.startswith("word/footer") for part in hidden_parts)


def test_docx_adapter_rejects_corrupt_docx_with_domain_error(tmp_path: Path) -> None:
    source = tmp_path / "broken.docx"
    source.write_bytes(b"not an OOXML package")

    with pytest.raises(InputValidationError, match="invalid DOCX") as error:
        DocxAdapter().load(source, render=False, workdir=tmp_path / "work")

    assert str(source) in str(error.value)


def test_docx_render_without_libreoffice_keeps_semantics_and_adds_warning(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = make_docx(
        tmp_path / "sample.docx",
        heading="Q2 Results",
        paragraphs=["Revenue increased."],
        rows=[["Region", "Total"]],
    )
    monkeypatch.setattr(docx_module, "find_libreoffice", lambda: None, raising=False)

    snapshot = DocxAdapter().load(source, render=True, workdir=tmp_path / "work")

    assert [block.text for block in snapshot.blocks[:2]] == [
        "Q2 Results",
        "Revenue increased.",
    ]
    assert snapshot.pages == []
    assert snapshot.page_count is None
    assert len(snapshot.warnings) == 1
    assert "LibreOffice was not found" in snapshot.warnings[0]


def test_docx_render_conversion_failure_keeps_semantics_and_adds_warning(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = make_docx(
        tmp_path / "sample.docx",
        heading="Q2 Results",
        paragraphs=["Revenue increased."],
        rows=[["Region", "Total"]],
    )
    executable = tmp_path / "soffice"
    monkeypatch.setattr(docx_module, "find_libreoffice", lambda: executable)

    def fail_conversion(*args: object, **kwargs: object) -> Path:
        raise RenderUnavailableError("LibreOffice timed out")

    monkeypatch.setattr(docx_module, "convert_docx_to_pdf", fail_conversion, raising=False)

    snapshot = DocxAdapter().load(source, render=True, workdir=tmp_path / "work")

    assert snapshot.blocks[0].text == "Q2 Results"
    assert snapshot.pages == []
    assert snapshot.warnings == ["DOCX visual rendering unavailable: LibreOffice timed out"]


def test_docx_render_reuses_pdf_adapter_pages_and_retains_logical_blocks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = make_docx(
        tmp_path / "sample.docx",
        heading="Q2 Results",
        paragraphs=["Revenue increased."],
        rows=[["Region", "Total"]],
    )
    converted = make_pdf(
        tmp_path / "sample.pdf",
        [[(72, 720, "Revenue 100")], [(72, 720, "Notes")]],
    )
    monkeypatch.setattr(docx_module, "find_libreoffice", lambda: tmp_path / "soffice")
    monkeypatch.setattr(
        docx_module,
        "convert_docx_to_pdf",
        lambda source, output_dir, executable: converted,
    )

    snapshot = DocxAdapter().load(source, render=True, workdir=tmp_path / "work")

    assert snapshot.page_count == 2
    assert [page.text for page in snapshot.pages] == ["Revenue 100", "Notes"]
    assert all(Path(page.render_path).is_file() for page in snapshot.pages)
    assert [block.content_type for block in snapshot.blocks[:3]] == [
        ContentType.HEADING,
        ContentType.PARAGRAPH,
        ContentType.TABLE,
    ]
    assert snapshot.warnings == [
        "DOCX rendered evidence alignment was incomplete"
    ]


def test_docx_render_links_unique_normalized_text_to_rendered_geometry(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = make_docx(
        tmp_path / "contract.docx",
        heading="第四条　付款条件",
        paragraphs=["发票开具后 30 天内付款"],
        rows=[["", ""]],
    )
    rendered = _rendered_snapshot("第四条 付款条件", "发票开具后 30 天内付款")
    monkeypatch.setattr(docx_module, "find_libreoffice", lambda: tmp_path / "soffice")
    monkeypatch.setattr(
        docx_module,
        "convert_docx_to_pdf",
        lambda source, output_dir, executable: tmp_path / "rendered.pdf",
    )
    monkeypatch.setattr(
        "artifactdiff.formats.docx.PdfAdapter.load", lambda *args, **kwargs: rendered
    )

    snapshot = DocxAdapter().load(source, render=True, workdir=tmp_path / "work")

    assert snapshot.blocks[0].metadata["rendered_page_index"] == 0
    assert snapshot.blocks[0].metadata["rendered_bbox"] == {
        "x0": 10.0,
        "y0": 20.0,
        "x1": 200.0,
        "y1": 35.0,
    }
    assert snapshot.blocks[1].metadata["rendered_bbox"] == {
        "x0": 10.0,
        "y0": 40.0,
        "x1": 200.0,
        "y1": 55.0,
    }
    assert snapshot.warnings == []


def test_docx_render_links_unique_near_exact_text_at_similarity_boundary(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    logical = "a" * 98 + "bc"
    rendered = _rendered_snapshot("a" * 98 + "de")
    source = make_docx(
        tmp_path / "contract.docx", heading=logical, paragraphs=[], rows=[["", ""]]
    )
    monkeypatch.setattr(docx_module, "find_libreoffice", lambda: tmp_path / "soffice")
    monkeypatch.setattr(
        docx_module,
        "convert_docx_to_pdf",
        lambda source, output_dir, executable: tmp_path / "rendered.pdf",
    )
    monkeypatch.setattr(
        "artifactdiff.formats.docx.PdfAdapter.load", lambda *args, **kwargs: rendered
    )

    snapshot = DocxAdapter().load(source, render=True, workdir=tmp_path / "work")

    assert snapshot.blocks[0].metadata["rendered_bbox"] == {
        "x0": 10.0,
        "y0": 20.0,
        "x1": 200.0,
        "y1": 35.0,
    }
    assert snapshot.warnings == []


def test_docx_render_leaves_duplicate_alignment_unset_and_warns_once(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = make_docx(
        tmp_path / "contract.docx",
        heading="Payment",
        paragraphs=["Payment"],
        rows=[["", ""]],
    )
    rendered = _rendered_snapshot("Payment", "Payment", "unmatched rendered line")
    monkeypatch.setattr(docx_module, "find_libreoffice", lambda: tmp_path / "soffice")
    monkeypatch.setattr(
        docx_module,
        "convert_docx_to_pdf",
        lambda source, output_dir, executable: tmp_path / "rendered.pdf",
    )
    monkeypatch.setattr(
        "artifactdiff.formats.docx.PdfAdapter.load", lambda *args, **kwargs: rendered
    )

    snapshot = DocxAdapter().load(source, render=True, workdir=tmp_path / "work")

    assert "rendered_bbox" not in snapshot.blocks[0].metadata
    assert "rendered_bbox" not in snapshot.blocks[1].metadata
    assert snapshot.warnings == [
        "DOCX rendered evidence alignment was incomplete"
    ]
    assert "Payment" not in snapshot.warnings[0]


def test_docx_render_output_directory_error_keeps_semantics_and_adds_one_warning(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = make_docx(
        tmp_path / "sample.docx",
        heading="Q2 Results",
        paragraphs=["Revenue increased."],
        rows=[["Region", "Total"]],
    )
    workdir = tmp_path / "render-work"
    blocked_output_dir = workdir / "libreoffice"
    original_mkdir = Path.mkdir

    def fail_output_mkdir(path: Path, *args: object, **kwargs: object) -> None:
        if path == blocked_output_dir:
            raise PermissionError("read-only render workdir")
        original_mkdir(path, *args, **kwargs)

    monkeypatch.setattr(Path, "mkdir", fail_output_mkdir)
    monkeypatch.setattr(docx_module, "find_libreoffice", lambda: tmp_path / "soffice")

    snapshot = DocxAdapter().load(source, render=True, workdir=workdir)

    assert [block.text for block in snapshot.blocks[:2]] == [
        "Q2 Results",
        "Revenue increased.",
    ]
    assert snapshot.pages == []
    assert snapshot.page_count is None
    assert len(snapshot.warnings) == 1
    assert "read-only render workdir" in snapshot.warnings[0]


def test_docx_render_invalid_converted_pdf_keeps_semantics(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = make_docx(
        tmp_path / "sample.docx",
        heading="Q2 Results",
        paragraphs=["Revenue increased."],
        rows=[["Region", "Total"]],
    )
    converted = tmp_path / "sample.pdf"
    converted.write_bytes(b"not a PDF")
    monkeypatch.setattr(docx_module, "find_libreoffice", lambda: tmp_path / "soffice")
    monkeypatch.setattr(
        docx_module,
        "convert_docx_to_pdf",
        lambda source, output_dir, executable: converted,
    )

    snapshot = DocxAdapter().load(source, render=True, workdir=tmp_path / "work")

    assert snapshot.blocks[0].text == "Q2 Results"
    assert snapshot.pages == []
    assert len(snapshot.warnings) == 1
    assert "invalid PDF" in snapshot.warnings[0]
