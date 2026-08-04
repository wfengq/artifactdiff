"""DOCX logical structure extraction with optional visual rendering."""

from __future__ import annotations

from collections.abc import Iterator
from difflib import SequenceMatcher
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.document import Document as DocumentType
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml.etree import XMLSyntaxError
from pydantic import JsonValue

from artifactdiff.alignment import align_sequences
from artifactdiff.contract.features import (
    CONTRACT_VISIBLE_TEXT_METADATA_KEY,
    FeatureInspectionError,
    inspect_docx_features,
    legacy_ooxml_paragraph_text,
    legacy_ooxml_table_text,
    visible_ooxml_paragraph_text,
    visible_ooxml_table_text,
)
from artifactdiff.errors import ArtifactDiffError, InputValidationError, RenderUnavailableError
from artifactdiff.formats.pdf import PdfAdapter
from artifactdiff.libreoffice import convert_docx_to_pdf, find_libreoffice
from artifactdiff.models import ContentBlock, ContentType, DocumentSnapshot, PageSnapshot
from artifactdiff.normalize import fingerprint, normalize_text

RENDERED_EVIDENCE_WARNING = "DOCX rendered evidence alignment was incomplete"
RENDERED_EVIDENCE_MIN_SIMILARITY = 0.98


def _rendered_text_similarity(before: ContentBlock, after: ContentBlock) -> float:
    if before.normalized_text == after.normalized_text:
        return 1.0
    return SequenceMatcher(
        None, before.normalized_text, after.normalized_text, autojunk=False
    ).ratio()


def _link_rendered_evidence(
    logical: list[ContentBlock], rendered: list[ContentBlock]
) -> bool:
    """Link unique near-exact logical/rendered text pairs and report incompleteness."""
    pairs = align_sequences(
        logical,
        rendered,
        key=lambda block: block.normalized_text,
        threshold=RENDERED_EVIDENCE_MIN_SIMILARITY,
    )
    eligible = [
        (before_index, after_index)
        for before_index, before in enumerate(logical)
        for after_index, after in enumerate(rendered)
        if _rendered_text_similarity(before, after) >= RENDERED_EVIDENCE_MIN_SIMILARITY
    ]
    logical_counts = {
        index: sum(before_index == index for before_index, _ in eligible)
        for index in range(len(logical))
    }
    rendered_counts = {
        index: sum(after_index == index for _, after_index in eligible)
        for index in range(len(rendered))
    }
    logical_indexes = {id(block): index for index, block in enumerate(logical)}
    rendered_indexes = {id(block): index for index, block in enumerate(rendered)}
    incomplete = False
    for pair in pairs:
        if pair.before is None or pair.after is None:
            incomplete = True
            continue
        before_index = logical_indexes[id(pair.before)]
        after_index = rendered_indexes[id(pair.after)]
        if (
            pair.after.page_index is None
            or pair.after.bbox is None
            or _rendered_text_similarity(pair.before, pair.after)
            < RENDERED_EVIDENCE_MIN_SIMILARITY
            or logical_counts[before_index] != 1
            or rendered_counts[after_index] != 1
        ):
            incomplete = True
            continue
        pair.before.metadata["rendered_page_index"] = pair.after.page_index
        pair.before.metadata["rendered_bbox"] = pair.after.bbox.model_dump(mode="json")
    return incomplete


def iter_body_items(document: DocumentType) -> Iterator[Paragraph | Table]:
    """Yield body paragraphs and tables in their original document order."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


class DocxAdapter:
    """Create a semantic snapshot from DOCX body structure."""

    def load(
        self, path: Path, *, render: bool, workdir: Path, force: bool = False
    ) -> DocumentSnapshot:
        """Extract ordered paragraphs and tables from a DOCX source."""
        try:
            document = Document(str(path))
        except (
            BadZipFile,
            KeyError,
            OSError,
            PackageNotFoundError,
            ValueError,
            XMLSyntaxError,
        ) as error:
            raise InputValidationError(f"invalid DOCX: {path}") from error
        blocks: list[ContentBlock] = []
        visible_text_by_block: dict[str, str] = {}
        for item in iter_body_items(document):
            block = self._body_block(item, len(blocks))
            if block is not None:
                blocks.append(block)
                visible_text_by_block[block.id] = self._body_visible_text(item)
        self._append_section_blocks(document, blocks, visible_text_by_block)
        try:
            features = inspect_docx_features(path)
        except FeatureInspectionError as error:
            raise InputValidationError(f"invalid DOCX: {path}") from error
        warnings: list[str] = []
        pages: list[PageSnapshot] = []
        if render:
            executable = find_libreoffice()
            if executable is None:
                warnings.append("DOCX visual rendering unavailable: LibreOffice was not found")
            else:
                try:
                    converted = convert_docx_to_pdf(path, workdir / "libreoffice", executable)
                except RenderUnavailableError as error:
                    warnings.append(f"DOCX visual rendering unavailable: {error}")
                else:
                    try:
                        rendered = PdfAdapter().load(
                            converted,
                            render=True,
                            workdir=workdir / "pdf-render",
                            force=force,
                        )
                    except ArtifactDiffError as error:
                        warnings.append(f"DOCX visual rendering unavailable: {error}")
                    else:
                        pages = rendered.pages
                        visible_blocks = [
                            block.model_copy(
                                update={
                                    "text": visible_text,
                                    "normalized_text": normalize_text(visible_text),
                                }
                            )
                            for block in blocks
                            if normalize_text(
                                visible_text := visible_text_by_block.get(block.id, block.text)
                            )
                        ]
                        if _link_rendered_evidence(visible_blocks, rendered.blocks):
                            warnings.append(RENDERED_EVIDENCE_WARNING)
        return DocumentSnapshot.from_path(
            path,
            pages=pages,
            blocks=blocks,
            warnings=warnings,
            metadata={
                CONTRACT_VISIBLE_TEXT_METADATA_KEY: {
                    block_id: text for block_id, text in visible_text_by_block.items()
                },
                "document_features": [feature.model_dump(mode="json") for feature in features],
            },
        )

    @staticmethod
    def _append_section_blocks(
        document: DocumentType,
        blocks: list[ContentBlock],
        visible_text_by_block: dict[str, str],
    ) -> None:
        for content_type, part_name in (
            (ContentType.HEADER, "header"),
            (ContentType.FOOTER, "footer"),
        ):
            seen: set[str] = set()
            for section in document.sections:
                part = getattr(section, part_name)
                text = "\n".join(
                    legacy_ooxml_paragraph_text(paragraph._p).strip()
                    for paragraph in part.paragraphs
                    if legacy_ooxml_paragraph_text(paragraph._p).strip()
                )
                if not text or text in seen:
                    continue
                seen.add(text)
                ordinal = len(blocks)
                block = ContentBlock(
                    id=f"docx:{ordinal}:{content_type}:{fingerprint(text)}",
                    ordinal=ordinal,
                    content_type=content_type,
                    text=text,
                    normalized_text=normalize_text(text),
                )
                blocks.append(block)
                visible_text_by_block[block.id] = "\n".join(
                    visible_ooxml_paragraph_text(paragraph._p).strip()
                    for paragraph in part.paragraphs
                    if visible_ooxml_paragraph_text(paragraph._p).strip()
                )

    @staticmethod
    def _body_block(item: Paragraph | Table, ordinal: int) -> ContentBlock | None:
        if isinstance(item, Paragraph):
            text = legacy_ooxml_paragraph_text(item._p).strip()
            if not text:
                return None
            style_name = item.style.name if item.style is not None else ""
            metadata: dict[str, JsonValue] = {}
            content_type = ContentType.PARAGRAPH
            if style_name.startswith("Heading "):
                try:
                    metadata["level"] = int(style_name.removeprefix("Heading "))
                except ValueError:
                    pass
                else:
                    content_type = ContentType.HEADING
            else:
                for prefix, kind in (("List Bullet", "bullet"), ("List Number", "number")):
                    if not style_name.startswith(prefix):
                        continue
                    suffix = style_name.removeprefix(prefix).strip()
                    metadata = {
                        "list_kind": kind,
                        "level": int(suffix) if suffix.isdigit() else 1,
                    }
                    break
        else:
            text = legacy_ooxml_table_text(item._tbl)
            if not text.strip():
                return None
            content_type = ContentType.TABLE
            metadata = {"rows": len(item.rows), "columns": len(item.columns)}

        return ContentBlock(
            id=f"docx:{ordinal}:{content_type}:{fingerprint(text)}",
            ordinal=ordinal,
            content_type=content_type,
            text=text,
            normalized_text=normalize_text(text),
            metadata=metadata,
        )

    @staticmethod
    def _body_visible_text(item: Paragraph | Table) -> str:
        if isinstance(item, Paragraph):
            return visible_ooxml_paragraph_text(item._p).strip()
        return visible_ooxml_table_text(item._tbl)
